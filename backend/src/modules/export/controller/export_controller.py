import uuid
from datetime import date

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response

from src.core.repositories.repositories import Repositories
from src.modules.auth.auth import get_current_user
from src.modules.export.schema.export_schemas import ExportFormat, ExportScope
from src.modules.export.service.export_service import ExportService
from src.modules.user.model.user_model import User
from src.modules.user.model.role.user_role import UserRole

router = APIRouter(prefix="/export", tags=["Export"])

CONTENT_TYPES = {
    ExportFormat.csv:  "text/csv",
    ExportFormat.xlsx: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def get_export_service() -> ExportService:
    return ExportService(Repositories())


@router.get(
    "/metrics/{researcher_id}",
    summary="Export researcher metrics.",
)
async def export_metrics(
        researcher_id: uuid.UUID,
        format:  ExportFormat  = Query(ExportFormat.xlsx),
        scope:   ExportScope   = Query(ExportScope.history),
        start_date: date | None = Query(None),
        end_date: date | None = Query(None),
        current_user: User     = Depends(get_current_user),
        service: ExportService = Depends(get_export_service),
) -> Response:
    if current_user.role != UserRole.admin and str(current_user.id) != str(researcher_id):
        from fastapi import HTTPException
        from starlette import status
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You can only export your own metrics.",
        )

    data, filename = await service.export(researcher_id, format, scope, start_date, end_date)

    return Response(
        content=data,
        media_type=CONTENT_TYPES[format],
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


@router.get(
    "/metrics",
    summary="Export all researcher metrics",
)
async def export_all_metrics(
        format: ExportFormat  = Query(ExportFormat.xlsx),
        scope:  ExportScope   = Query(ExportScope.history),
        start_date: date | None = Query(None),
        end_date: date | None = Query(None),
        current_user: User    = Depends(get_current_user),
        service: ExportService = Depends(get_export_service),
) -> Response:
    from fastapi import HTTPException
    from starlette import status

    if current_user.role != UserRole.admin:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only admins can export all metrics.",
        )

    async with Repositories() as repos:
        users = await repos.users.get_all()

    all_data = b""
    all_rows = []

    from src.modules.export.schema.export_schemas import ExportFormat as EF
    from src.modules.export.service.export_service import HEADERS

    for user in users:
        if not user.active:
            continue

        try:
            rows = await service.build_rows(user.id, scope, start_date, end_date)
            all_rows.extend(rows)

        except Exception:
            continue

    if format == EF.csv:
        import csv, io
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(HEADERS)
        writer.writerows(all_rows)
        all_data = buf.getvalue().encode("utf-8-sig")
        filename = "metrics_all.csv"

    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "All Metrics"
        header_fill = PatternFill("solid", fgColor="1A3A5C")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        for col_idx, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_idx, row in enumerate(all_rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        ws.freeze_panes = "A2"
        import io
        buf = io.BytesIO()
        wb.save(buf)
        all_data = buf.getvalue()
        filename = "metrics_all.xlsx"

    return Response(
        content=all_data,
        media_type=CONTENT_TYPES[format],
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )