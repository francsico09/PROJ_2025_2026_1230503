import csv
import io
import uuid

from datetime import date, datetime, time, timezone

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from starlette import status

from src.core.repositories.repositories import Repositories
from src.modules.export.schema.export_schemas import ExportFormat, ExportScope

HEADERS = [
    "Researcher Name",
    "Email",
    "ORCID iD",
    "Scholar ID",
    "Date",
    "Source",
    "Source URL",
    "h-index",
    "i10-index",
    "Total Citations",
    "Total Publications",
]

class ExportService:

    def __init__(self, repos: Repositories) -> None:
        self._repos = repos

    async def export(
            self,
            researcher_id: uuid.UUID,
            fmt: ExportFormat,
            scope: ExportScope,
            start_date: date | None = None,
            end_date: date | None = None,
    ) -> tuple[bytes, str]:
        """
        Gera o ficheiro de exportação em memória.

        :param researcher_id: id do investigador
        :param fmt: ExportFormat.csv ou ExportFormat.xlsx
        :param scope: ExportScope.latest, ExportScope.history, ou ExportScope.custom
        :param start_date: data de início (apenas para scope custom)
        :param end_date: data de fim (apenas para scope custom)
        :return: (bytes do ficheiro, filename)
        :raises HTTPException 404: se o investigador não existir
        """
        rows = await self.build_rows(researcher_id, scope, start_date, end_date)

        if fmt == ExportFormat.csv:
            data, filename = self._to_csv(rows, researcher_id)
        else:
            data, filename = self._to_xlsx(rows, researcher_id)

        return data, filename

    # Data fetching
    async def build_rows(
            self,
            researcher_id: uuid.UUID,
            scope: ExportScope,
            start_date: datetime | None = None,
            end_date: datetime | None = None,
    ) -> list[list]:
        async with self._repos as repos:
            user = await repos.users.get_by_id(researcher_id)

            if not user:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail="Researcher not found",
                )

            profile = user.researcherProfile
            orcid_id   = profile.orcid      if profile else ""
            scholar_id = profile.scholar_id if profile else ""

            metrics = await repos.metrics.get_by_researcher_id(profile.id)

        if not metrics:
            return []

        # Filter by date range if scope is custom
        if scope == ExportScope.custom:
            if start_date and end_date:
                start_date = datetime.combine(start_date, time.min).replace(tzinfo=timezone.utc)
                end_date = datetime.combine(end_date, time.max).replace(tzinfo=timezone.utc)

                metrics = [
                    m for m in metrics
                    if m.date and start_date <= m.date <= end_date
                ]

            elif start_date:
                start_date = datetime.combine(start_date, time.min).replace(tzinfo=timezone.utc)
                metrics = [
                    m for m in metrics
                    if m.date and m.date >= start_date
                ]

            elif end_date:
                end_date = datetime.combine(end_date, time.max).replace(tzinfo=timezone.utc)
                metrics = [
                    m for m in metrics
                    if m.date and m.date <= end_date
                ]

        if scope == ExportScope.latest:
            latest: dict[str, object] = {}
            for m in metrics:
                source_name = m.source.name if m.source else "unknown"
                existing = latest.get(source_name)
                if not existing or m.date > existing.date:
                    latest[source_name] = m
            metrics = list(latest.values())

        metrics = sorted(metrics, key=lambda m: m.date or date.min, reverse=True)

        rows = []
        for m in metrics:
            rows.append([
                user.name,
                user.email,
                orcid_id,
                scholar_id,
                str(m.date) if m.date else "",
                m.source.name if m.source else "",
                m.source.url  if m.source else "",
                m.h_index,
                m.i10_index,
                m.total_citations,
                m.total_publications,
            ])

        return rows

    # Formatters
    @staticmethod
    def _to_csv(rows: list[list], researcher_id: uuid.UUID) -> tuple[bytes, str]:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(HEADERS)
        writer.writerows(rows)
        data = buf.getvalue().encode("utf-8-sig")
        filename = f"metrics_{researcher_id}.csv"
        return data, filename

    @staticmethod
    def _to_xlsx(rows: list[list], researcher_id: uuid.UUID) -> tuple[bytes, str]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Metrics"

        # Header row styling
        header_fill = PatternFill("solid", fgColor="1A3A5C")
        header_font = Font(color="FFFFFF", bold=True, size=11)

        for col_idx, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

        # Freeze header row
        ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        filename = f"metrics_{researcher_id}.xlsx"
        return data, filename