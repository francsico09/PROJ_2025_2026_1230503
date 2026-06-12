import io
import uuid
import pytest
from datetime import datetime, timezone, timedelta, date
from unittest.mock import AsyncMock, MagicMock
from fastapi import HTTPException
from openpyxl import load_workbook

from src.modules.export.service.export_service import ExportService, HEADERS
from src.core.domain.export.export_schema.export_schemas import ExportFormat, ExportScope
from src.core.domain.researcher_metric.researcher_metric_model.source_model import SourceName
from src.core.domain.researcher_metric.researcher_metric_model.source_model import Source

# Fixtures

@pytest.fixture
def mock_repos():
    """Cria um mock das repositories com contexto assíncrono."""
    repos = AsyncMock()
    repos.metrics = AsyncMock()
    repos.profiles = AsyncMock()
    repos.users = AsyncMock()
    repos.extraction_runs = AsyncMock()
    repos.commit = AsyncMock()
    repos.__aenter__ = AsyncMock(return_value=repos)
    repos.__aexit__ = AsyncMock(return_value=None)
    return repos

# =============================================================================
# Helper classes/mocks
# =============================================================================

class MockSource:
    def __init__(self, name=SourceName.scholar, url="http://scholar.com"):
        self.name = name
        self.url = url

class MockMetric:
    def __init__(self, dt, source: Source = None):
        self.date = dt

        if source:
            self.source = source
        else:
            self.source = MockSource()

        self.h_index = 10
        self.i10_index = 10
        self.total_citations = 100
        self.total_publications = 50


# =============================================================================
# build_rows — Core Data Fetching and Aggregation
# =============================================================================

class TestExportServiceBuildRows:
    """Tests for fetching, filtering and aggregating user metrics."""

    @pytest.mark.asyncio
    async def test_build_rows_user_not_found(self, mock_repos):
        """When a user does not exist, an HTTPException (404) should be raised."""
        # Arrange
        mock_repos.users.get_by_id = AsyncMock(return_value=None)
        service = ExportService(mock_repos)

        # Act & Assert
        with pytest.raises(HTTPException) as exc:
            await service.build_rows(uuid.uuid4(), ExportScope.history)

        assert exc.value.status_code == 404

    @pytest.mark.asyncio
    async def test_build_rows_history_scope(self, mock_repos, sample_user, sample_profile):
        """When scope is history, all metrics are returned sorted by date descending."""
        # Arrange
        sample_user.researcherProfile = sample_profile
        mock_repos.users.get_by_id = AsyncMock(return_value=sample_user)

        now = datetime.now(timezone.utc)
        metrics = [
            MockMetric(now - timedelta(days=2)),
            MockMetric(now),
            MockMetric(now - timedelta(days=5)),
        ]

        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=metrics)
        service = ExportService(mock_repos)

        # Act
        rows = await service.build_rows(sample_user.id, ExportScope.history)

        # Assert
        assert len(rows) == 3
        assert str(now) in rows[0][4]
        assert str(now - timedelta(days=5)) in rows[2][4]

    @pytest.mark.asyncio
    async def test_build_rows_latest_scope(self, mock_repos, sample_user, sample_profile):
        """When scope is latest, only the most recent metric per source is returned."""
        # Arrange
        sample_user.researcherProfile = sample_profile
        mock_repos.users.get_by_id = AsyncMock(return_value=sample_user)

        now = datetime.now(timezone.utc)
        metrics = [
            MockMetric(now - timedelta(days=10), source=Source(SourceName.scholar, "")),
            MockMetric(now, source=Source(SourceName.scholar, "")),
            MockMetric(now - timedelta(days=2), source=Source(SourceName.wos, "")),
        ]
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=metrics)
        service = ExportService(mock_repos)

        # Act
        rows = await service.build_rows(sample_user.id, ExportScope.latest)

        # Assert
        assert len(rows) == 2
        sources = [row[5] for row in rows]
        assert SourceName.scholar in sources and SourceName.wos in sources


    @pytest.mark.asyncio
    async def test_build_rows_custom_scope_both_dates(self, mock_repos, sample_user, sample_profile):
        """When scope is custom with start and end dates, metrics are strictly bounded."""
        # Arrange
        sample_user.researcherProfile = sample_profile
        mock_repos.users.get_by_id = AsyncMock(return_value=sample_user)

        base_date = datetime(2023, 1, 15, tzinfo=timezone.utc)
        metrics = [
            MockMetric(base_date - timedelta(days=20)), # Before start
            MockMetric(base_date),                      # Inside range
            MockMetric(base_date + timedelta(days=20)), # After end
        ]
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=metrics)
        service = ExportService(mock_repos)

        # Act
        start = date(2023, 1, 1)
        end = date(2023, 1, 31)
        rows = await service.build_rows(sample_user.id, ExportScope.custom, start_date=start, end_date=end)

        # Assert
        assert len(rows) == 1
        assert str(base_date) in rows[0][4]


# =============================================================================
# export — Individual Researcher Export (Edge Cases)
# =============================================================================

class TestExportServiceIndividual:
    """Tests for individual export, focusing on edge cases and user states."""

    @pytest.mark.asyncio
    async def test_export_user_not_found(self, mock_repos):
        """When the requested user ID does not exist, an HTTP 404 is raised."""
        # Arrange
        mock_repos.users.get_by_id = AsyncMock(return_value=None)
        service = ExportService(mock_repos)

        # Act & Assert
        with pytest.raises(HTTPException) as exc:
            await service.export(uuid.uuid4(), ExportFormat.csv, ExportScope.latest)

        assert exc.value.status_code == 404
        assert exc.value.detail == "Researcher not found"

    @pytest.mark.asyncio
    async def test_export_user_inactive(self, mock_repos, sample_user):
        """When the user is found but is inactive, the export should be denied."""
        # Arrange
        sample_user.active = False
        mock_repos.users.get_by_id = AsyncMock(return_value=sample_user)
        service = ExportService(mock_repos)

        # Act & Assert
        with pytest.raises(HTTPException) as exc:
            await service.export(sample_user.id, ExportFormat.xlsx, ExportScope.latest)

        assert exc.value.status_code == 400

    @pytest.mark.asyncio
    async def test_export_individual_csv(self, mock_repos, sample_user, sample_profile):
        """CSV export should return bytes containing headers and researcher data."""
        # Arrange
        sample_user.researcherProfile = sample_profile
        mock_repos.users.get_by_id = AsyncMock(return_value=sample_user)
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=[MockMetric(datetime.now(timezone.utc))])
        service = ExportService(mock_repos)

        # Act
        file_bytes, filename = await service.export(sample_user.id, ExportFormat.csv, ExportScope.history)

        # Assert
        assert filename.endswith(".csv")
        assert isinstance(file_bytes, bytes)
        content = file_bytes.decode("utf-8-sig")
        assert HEADERS[0] in content
        assert sample_user.name in content

    @pytest.mark.asyncio
    async def test_export_individual_xlsx(self, mock_repos, sample_user, sample_profile):
        """XLSX export should return a valid Excel byte stream with the correct worksheet."""
        # Arrange
        sample_user.researcherProfile = sample_profile
        mock_repos.users.get_by_id = AsyncMock(return_value=sample_user)
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=[MockMetric(datetime.now(timezone.utc))])
        service = ExportService(mock_repos)

        # Act
        file_bytes, filename = await service.export(sample_user.id, ExportFormat.xlsx, ExportScope.history)

        # Assert
        assert filename.endswith(".xlsx")
        wb = load_workbook(io.BytesIO(file_bytes))
        ws = wb.active
        assert ws.title == "Metrics"
        assert ws.cell(row=1, column=1).value == HEADERS[0]
        assert ws.cell(row=2, column=1).value == sample_user.name


# =============================================================================
# export_all — Bulk Export Scenarios
# =============================================================================

class TestExportServiceBulk:
    """Tests for bulk exporting, ensuring correct behavior across all dataset variations."""

    @pytest.mark.asyncio
    async def test_export_all_complete_success(self, mock_repos, sample_profile):
        """When all users are active and fetch successfully, all data is exported."""
        # Arrange
        user1 = MagicMock(id=uuid.uuid4(), active=True, name="User One", researcherProfile=sample_profile)
        user2 = MagicMock(id=uuid.uuid4(), active=True, name="User Two", researcherProfile=sample_profile)

        mock_repos.users.get_all = AsyncMock(return_value=[user1, user2])

        async def side_effect_get_user(uid):
            return user1 if uid == user1.id else user2

        mock_repos.users.get_by_id = AsyncMock(side_effect=side_effect_get_user)
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=[MockMetric(datetime.now(timezone.utc))])

        service = ExportService(mock_repos)

        # Act
        file_bytes, filename = await service.export_all(ExportFormat.csv, ExportScope.latest)

        # Assert
        content = file_bytes.decode("utf-8-sig")
        lines = [line for line in content.split("\n") if line.strip()]

        assert filename == "metrics_all.csv"
        assert len(lines) == 3  # 1 Header + 2 Users
        assert "User One" in content
        assert "User Two" in content

    @pytest.mark.asyncio
    async def test_export_all_complete_failure(self, mock_repos):
        """When no users are active, the export returns a file with only headers."""
        # Arrange
        inactive1 = MagicMock(id=uuid.uuid4(), active=False)
        inactive2 = MagicMock(id=uuid.uuid4(), active=False)

        mock_repos.users.get_all = AsyncMock(return_value=[inactive1, inactive2])
        service = ExportService(mock_repos)

        # Act
        file_bytes, filename = await service.export_all(ExportFormat.csv, ExportScope.latest)

        # Assert
        content = file_bytes.decode("utf-8-sig")
        lines = [line for line in content.split("\n") if line.strip()]

        assert filename == "metrics_all.csv"
        assert len(lines) == 1  # Only the header row should be present
        assert HEADERS[0] in lines[0]
        # Ensure build_rows was never called since all were inactive
        mock_repos.users.get_by_id.assert_not_called()

    @pytest.mark.asyncio
    async def test_export_all_partial_success(self, mock_repos, sample_profile):
        """When some users fail or are inactive, only the valid ones are exported."""
        # Arrange
        active_user = MagicMock(id=uuid.uuid4(), active=True, name="Active User", researcherProfile=sample_profile)
        inactive_user = MagicMock(id=uuid.uuid4(), active=False, name="Inactive User")
        failing_user = MagicMock(id=uuid.uuid4(), active=True, name="Failing User")

        mock_repos.users.get_all = AsyncMock(return_value=[active_user, inactive_user, failing_user])

        # active_user succeeds, failing_user raises 404 (simulating DB desync)
        async def side_effect_get_user(uid):
            if uid == active_user.id:
                return active_user
            elif uid == failing_user.id:
                raise HTTPException(status_code=404)
            return None

        mock_repos.users.get_by_id = AsyncMock(side_effect=side_effect_get_user)
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=[MockMetric(datetime.now(timezone.utc))])

        service = ExportService(mock_repos)

        # Act
        file_bytes, filename = await service.export_all(ExportFormat.csv, ExportScope.latest)

        # Assert
        content = file_bytes.decode("utf-8-sig")
        lines = [line for line in content.split("\n") if line.strip()]

        assert len(lines) == 2  # 1 Header + 1 Active User
        assert "Active User" in content
        assert "Inactive User" not in content
        assert "Failing User" not in content


    @pytest.mark.asyncio
    async def test_export_all_skips_inactive_and_failing_users(self, mock_repos, sample_user, sample_profile):
        """Users that are inactive or raise exceptions during row building should be skipped gracefully."""
        # Arrange
        active_user = sample_user
        active_user.active = True
        active_user.researcherProfile = sample_profile
        active_user.name = "Active USer"

        inactive_user = MagicMock()
        inactive_user.active = False
        inactive_user.name = "Inactive User"

        failing_user = MagicMock()
        failing_user.active = True
        failing_user.id = uuid.uuid4()
        failing_user.name = "Failing User"

        mock_repos.users.get_all = AsyncMock(return_value=[inactive_user, failing_user, active_user])

        # Define get_by_id to raise 404 for failing_user and return active_user
        async def side_effect_get_user(uid):
            if uid == failing_user.id:
                raise HTTPException(status_code=404)
            return active_user

        mock_repos.users.get_by_id = AsyncMock(side_effect=side_effect_get_user)
        mock_repos.metrics.get_by_researcher_id = AsyncMock(return_value=[MockMetric(datetime.now(timezone.utc))])

        # Needs to mock Repositories context manager usage if not fixed,
        # assuming you fix the DI bug mentioned, `service = ExportService(mock_repos)` will suffice.
        service = ExportService(mock_repos)

        # Act
        file_bytes, filename = await service.export_all(ExportFormat.csv, ExportScope.latest)

        # Assert
        content = file_bytes.decode("utf-8-sig")
        assert filename == "metrics_all.csv"
        lines = [line for line in content.split("\n") if line.strip()]
        assert len(lines) == 2
        assert active_user.name in content
        assert inactive_user.name not in content